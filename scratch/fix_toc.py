import sys
try:
    from docx import Document
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def fix_toc_v4(input_path, output_path):
    doc = Document(input_path)
    
    # We must evaluate doc.paragraphs ONCE to avoid dynamic shifting during deletion.
    paragraphs = list(doc.paragraphs)
    
    start_idx = -1
    end_idx = -1
    for i, p in enumerate(paragraphs[:200]):
        if "3.4 Data Flow" in p.text and "\t16" in p.text:
            start_idx = i
        if "Chapter 4\t24" in p.text:
            end_idx = i
            break
            
    if start_idx != -1 and end_idx != -1:
        toc_entries = [
            "3.4.1 Data Flow for Object Detection\t16",
            "3.4.2 Detection Result Flow\t18",
            "3.4.3 Sequence Diagram\t18",
            "3.4.4 Activity Diagram\t20",
            "3.4.5 System Overview Diagram\t21",
            "3.4.6 System Architecture Diagram\t22",
            "3.4.7 Use Case Diagram\t23"
        ]
        
        toc_3_style = paragraphs[start_idx].style 
        for i in range(start_idx - 10, start_idx):
            if paragraphs[i].text.startswith("3.3.1"):
                toc_3_style = paragraphs[i].style
                break
                
        # To delete the middle paragraphs completely from the XML tree:
        for i in range(start_idx + 1, end_idx):
            p = paragraphs[i]
            p_element = p._element
            p_element.getparent().remove(p_element)
            
        # Insert the correct entries before the end_idx paragraph
        # Wait, since end_idx was NOT deleted, we can still use paragraphs[end_idx]
        # because insert_paragraph_before operates on the underlying XML element
        insert_target = paragraphs[end_idx]
        for entry in toc_entries:
            new_p = insert_target.insert_paragraph_before(entry)
            new_p.style = toc_3_style
            
    doc.save(output_path)
    print("TOC Fixed Perfectly and Saved")

if __name__ == "__main__":
    fix_toc_v4(sys.argv[1], sys.argv[2])
