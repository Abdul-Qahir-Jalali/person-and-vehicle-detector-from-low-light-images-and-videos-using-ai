import sys
try:
    from docx import Document
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def fix_final_formatting(input_path, output_path):
    doc = Document(input_path)
    
    paragraphs = list(doc.paragraphs)
    
    found_fig = False
    found_undertaking = False
    found_ack = False
    
    undertaking_text = [
        'I certify that the research work titled "Robust Object Detection System Under Adverse Weather Conditions" is my own work. The work has not been presented elsewhere for assessment. Where material has been used from other sources it has been properly acknowledged/referred.',
        '',
        "(Student’s Signature)\t\t(Student’s Signature)",
        "Abdul Qahir Jalali\t\tRaja Sharyar Muneer",
        "2021-UMDB-001017\t\t2021-UMDB-001041"
    ]
    
    ack_text = [
        "All praises are for Allah, the Lord of all the worlds. It is only because of the boundless grace of Allah that we have been able to complete our project. Peace be upon Prophet Muhammad ﷺ whom Allah sent as a mercy to all the humanity.",
        "",
        "We extend our heartfelt gratitude to all the teachers who have guided us throughout our degree journey. In particular, we wish to thank our supervisor, Dr. Maryam Bibi for her invaluable support, guidance, and encouragement throughout the project development and documentation process. Her dedication and insight have been instrumental not only in this project but throughout our entire academic journey.",
        "",
        "We also wish to express our deep appreciation to our parents, family, and friends for their unwavering support and patience during our research and writing journey. Their prayers, encouragement, and belief in our abilities have been a constant source of strength and motivation, making this accomplishment possible."
    ]
    
    for i, p in enumerate(paragraphs):
        txt = p.text.strip().lower()
        
        # 1. Section 4.1.3 Figure 4.0.1 replacement
        if "test results" in txt and "figure 4" in txt and not found_fig:
            p.text = "[Placeholder: Graph illustrating the YOLO model's training and validation loss curves, showing steady convergence over 20 epochs, alongside the mAP score progression.]"
            for run in p.runs:
                run.bold = False
            found_fig = True
            
        if txt == "undertaking" and not found_undertaking:
            insert_idx = i + 1
            while insert_idx < len(paragraphs) and not paragraphs[insert_idx].text.strip():
                p_empty = paragraphs[insert_idx]._element
                parent = p_empty.getparent()
                if parent is not None:
                    parent.remove(p_empty)
                insert_idx += 1
                
            insert_target = paragraphs[insert_idx] if insert_idx < len(paragraphs) else doc.add_paragraph()
            for text in undertaking_text:
                new_p = insert_target.insert_paragraph_before(text)
            found_undertaking = True
            
        if txt == "acknowledgements" and not found_ack:
            insert_idx = i + 1
            while insert_idx < len(paragraphs) and not paragraphs[insert_idx].text.strip():
                p_empty = paragraphs[insert_idx]._element
                parent = p_empty.getparent()
                if parent is not None:
                    parent.remove(p_empty)
                insert_idx += 1
                
            insert_target = paragraphs[insert_idx] if insert_idx < len(paragraphs) else doc.add_paragraph()
            for text in ack_text:
                new_p = insert_target.insert_paragraph_before(text)
            found_ack = True
            
    for p in paragraphs:
        style_name = p.style.name.lower() if p.style else ""
        if "heading" in style_name or "toc" in style_name or "title" in style_name:
            continue
            
        for run in p.runs:
            if run.bold is True:
                run.bold = None

    doc.save(output_path)
    print(f"Final Formatting Fixes Applied. Fig: {found_fig}, Under: {found_undertaking}, Ack: {found_ack}")

if __name__ == "__main__":
    fix_final_formatting(sys.argv[1], sys.argv[2])
