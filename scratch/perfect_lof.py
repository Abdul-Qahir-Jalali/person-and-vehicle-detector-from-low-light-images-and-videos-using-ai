import sys
try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def fix_lof_perfectly(input_path, output_path):
    doc = Document(input_path)
    
    # 1. New figures to insert
    figures = [
        ("Figure 3.1: Sequence Diagram for Object Detection Flow", "18"),
        ("Figure 3.2: Activity Diagram of User Upload Process", "20"),
        ("Figure 3.3: High-level System Overview", "21"),
        ("Figure 3.4: Detailed System Architecture", "22"),
        ("Figure 3.5: Use Case Diagram for User Interactions", "23"),
        ("Figure 4.1: YOLO Training Validation Loss and mAP Score", "26"),
        ("Figure 5.1: Web App Home Page Interface", "31"),
        ("Figure 5.2: Detection Results Interface", "34")
    ]

    # 2. Identify LOF paragraphs
    lof_paras = []
    in_lof = False
    for p in doc.paragraphs:
        txt = p.text.strip().lower()
        if txt == "list of figures" and p.style.name == "Heading 1":
            in_lof = True
            continue
        if in_lof:
            if "table of figures" in p.style.name.lower():
                lof_paras.append(p)
            elif p.style.name == "Heading 1" or p.style.name == "Normal" and p.text.strip() == "Chapter 1":
                break
                
    # 3. Overwrite LOF paragraphs
    for i, (txt, pg) in enumerate(figures):
        if i < len(lof_paras):
            p = lof_paras[i]
            for r in p.runs:
                r.text = ""
                
            p.style = "table of figures"
            
            p.add_run(txt)
            r_tab = p.add_run()
            r_tab._element.append(OxmlElement('w:tab'))
            p.add_run(pg)
        else:
            prev_p = lof_paras[-1] if lof_paras else None
            if prev_p:
                new_p_xml = OxmlElement('w:p')
                pPr = OxmlElement("w:pPr")
                pStyle = OxmlElement("w:pStyle")
                pStyle.set(qn("w:val"), "TableofFigures") # Usually 'TableofFigures' or 'tableoffigures'
                pPr.append(pStyle)
                
                run_txt = OxmlElement('w:r')
                t_txt = OxmlElement('w:t')
                t_txt.text = txt
                run_txt.append(t_txt)
                
                run_tab = OxmlElement('w:r')
                run_tab.append(OxmlElement('w:tab'))
                
                run_pg = OxmlElement('w:r')
                t_pg = OxmlElement('w:t')
                t_pg.text = pg
                run_pg.append(t_pg)
                
                new_p_xml.append(pPr)
                new_p_xml.append(run_txt)
                new_p_xml.append(run_tab)
                new_p_xml.append(run_pg)
                
                last_p = lof_paras[-1]._element if lof_paras else None
                if last_p is not None:
                    last_p.addnext(new_p_xml)
                    lof_paras.append(new_p_xml) 

    # Remove leftover LOF paragraphs
    if len(lof_paras) > len(figures):
        for i in range(len(figures), len(lof_paras)):
            p = lof_paras[i]
            if hasattr(p, '_element'):
                parent = p._element.getparent()
                if parent is not None:
                    parent.remove(p._element)

    doc.save(output_path)
    print("Perfect LOF generated.")

if __name__ == "__main__":
    fix_lof_perfectly(sys.argv[1], sys.argv[2])
