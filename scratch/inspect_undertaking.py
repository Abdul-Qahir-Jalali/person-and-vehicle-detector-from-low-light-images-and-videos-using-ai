import sys
try:
    from docx import Document
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def inspect_undertaking(path):
    doc = Document(path)
    
    start = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip().lower()
        if txt == "undertaking":
            start = True
            
        if start:
            print(f"[{i}] PARA: {p.text}")
            
        if start and i > 50: # just limit
            break
            
    print("Checking tables for undertaking...")
    for t in doc.tables:
        text = "\n".join(cell.text for row in t.rows for cell in row.cells).lower()
        if "student’s signature" in text or "abdul qahir" in text:
            print(f"Found table with signatures!")
            for row_idx, row in enumerate(t.rows):
                for col_idx, cell in enumerate(row.cells):
                    print(f"  Row {row_idx}, Col {col_idx}: {cell.text.strip()}")

if __name__ == "__main__":
    inspect_undertaking(sys.argv[1])
