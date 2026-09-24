import sys
import os

try:
    from docx import Document
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def extract_docx(path):
    doc = Document(path)
    text = []
    for p in doc.paragraphs:
        if p.text.strip():
            text.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)
    return "\n".join(text)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(1)
    path = sys.argv[1]
    print(extract_docx(path))
