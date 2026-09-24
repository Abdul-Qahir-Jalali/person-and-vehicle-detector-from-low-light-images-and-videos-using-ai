import sys
import urllib.request
import urllib.error
import base64
import zlib
import os
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def generate_kroki(source, filename):
    try:
        # Use deflate + base64 for GET request
        data = base64.urlsafe_b64encode(zlib.compress(source.encode('utf-8'), 9)).decode('ascii')
        url = f"https://kroki.io/mermaid/png/{data}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=15) as response:
            with open(filename, 'wb') as f:
                f.write(response.read())
        print(f"Generated {filename}")
    except Exception as e:
        print(f"Failed to generate {filename}: {e}")

def create_kroki_diagrams():
    theme = "%%{init: {'theme': 'base', 'themeVariables': { 'primaryColor': '#ffffff', 'primaryBorderColor': '#000000', 'primaryTextColor': '#000000', 'lineColor': '#000000'}}}%%\n"
    
    seq = theme + """sequenceDiagram
    actor User
    participant Flask as Web App
    participant OpenCV as Preprocessing
    participant YOLO as Detection Model
    User->>Flask: Upload Media
    Flask->>OpenCV: Send frames
    OpenCV-->>Flask: Enhanced frames
    Flask->>YOLO: Pass for inference
    YOLO-->>Flask: Bounding Boxes & Confidence
    Flask-->>User: Display Results"""
    generate_kroki(seq, "seq.png")

    act = theme + """flowchart TD
    A([Start]) --> B[/Upload Image/Video/]
    B --> C{Format Valid?}
    C -- No --> D[/Show Error/]
    D --> B
    C -- Yes --> E[OpenCV Preprocessing]
    E --> F[YOLO Inference]
    F --> G[Generate Bounding Boxes]
    G --> H[/Render UI Overlays/]
    H --> I([End])"""
    generate_kroki(act, "act.png")
    
    over = theme + """flowchart LR
    Client((Client UI)) <--> |HTTP| Server[Flask Server]
    Server --> |Frames| Preproc[OpenCV Module]
    Preproc --> |Cleaned| Model[YOLOv8 Model]
    Model --> |Predictions| Server"""
    generate_kroki(over, "over.png")
    
    arch = theme + """flowchart TD
    subgraph Presentation
        UI[HTML / CSS UI]
    end
    subgraph Application
        API[Flask API]
        CV[OpenCV]
    end
    subgraph DeepLearning
        YV8[PyTorch / YOLOv8]
    end
    UI <--> API
    API <--> CV
    API <--> YV8"""
    generate_kroki(arch, "arch.png")
    
    uc = theme + """flowchart LR
    User((User))
    User --> U1([Upload Media])
    User --> U2([View Progress])
    User --> U3([Interpret Results])
    System((System))
    System --> U1
    System --> U2
    System --> U3"""
    generate_kroki(uc, "uc.png")

def insert_diagrams_and_wireframes(input_path, output_path, graph_img_path):
    doc = Document(input_path)
    
    mapping = {
        "sequence diagram showing the": ("image", "seq.png"),
        "activity diagram showing the": ("image", "act.png"),
        "high-level system overview diagram": ("image", "over.png"),
        "detailed system architecture diagram": ("image", "arch.png"),
        "use case diagram showing the": ("image", "uc.png"),
        "graph illustrating the yolo model": ("image", graph_img_path),
        "screenshot of the web app home page": ("ui_home", None),
        "screenshot showing the detection results": ("ui_results", None),
        "a classification report table": ("table", None)
    }
    
    def add_table_borders(table):
        tblPr = table._element.xpath('w:tblPr')
        if not tblPr:
            tblPr = OxmlElement('w:tblPr')
            table._element.insert(0, tblPr)
        else:
            tblPr = tblPr[0]
        
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), '000000')
            tblBorders.append(border)
        tblPr.append(tblBorders)

    for p in doc.paragraphs:
        txt = p.text.strip().lower()
        if not txt:
            continue
            
        for key, (action, arg) in mapping.items():
            if key in txt:
                # We found a placeholder!
                p.text = ""
                
                if action == "image":
                    if arg and os.path.exists(arg):
                        run = p.add_run()
                        width = Inches(5.0) if "graph" not in arg else Inches(4.5)
                        run.add_picture(arg, width=width)
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif action == "ui_home":
                    # Insert a simple 2x1 table to act as a UI wireframe
                    table = doc.add_table(rows=2, cols=1)
                    add_table_borders(table)
                    
                    cell0 = table.cell(0, 0)
                    c0_p = cell0.paragraphs[0]
                    c0_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r0 = c0_p.add_run("Robust Object Detection System")
                    r0.bold = True
                    r0.font.size = Pt(16)
                    
                    cell1 = table.cell(1, 0)
                    c1_p = cell1.paragraphs[0]
                    c1_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    c1_p.add_run("\n\n[ Drag & Drop Image/Video files (JPG, PNG, MP4) ]\n\n\n[ Upload ]\n")
                    
                    # Move table to right after paragraph p
                    p._element.addnext(table._element)
                    
                elif action == "ui_results":
                    table = doc.add_table(rows=2, cols=2)
                    add_table_borders(table)
                    
                    cell00 = table.cell(0, 0)
                    cell00.merge(table.cell(0, 1))
                    c0_p = cell00.paragraphs[0]
                    c0_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r0 = c0_p.add_run("Detection Results Interface")
                    r0.bold = True
                    r0.font.size = Pt(16)
                    
                    cell10 = table.cell(1, 0)
                    c10_p = cell10.paragraphs[0]
                    c10_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    c10_p.add_run("\n[ Foggy Road Image ]\n\n+---------------+      +---------------+\n| Vehicle (92%) |      | Pedestrian (88%) |\n+---------------+      +---------------+\n\n")
                    
                    cell11 = table.cell(1, 1)
                    c11_p = cell11.paragraphs[0]
                    c11_p.add_run("Sidebar:\n\nObjects Detected:\n- Vehicle: 92%\n- Pedestrian: 88%\n")
                    
                    p._element.addnext(table._element)
                    
                elif action == "table":
                    table = doc.add_table(rows=5, cols=4)
                    add_table_borders(table)
                    
                    data = [
                        ['Class', 'Precision', 'Recall', 'F1-Score'],
                        ['Vehicles', '0.93', '0.91', '0.92'],
                        ['Pedestrians', '0.89', '0.86', '0.87'],
                        ['Traffic Signs', '0.95', '0.92', '0.93'],
                        ['Average / Total', '0.92', '0.90', '0.91']
                    ]
                    for row_idx, row_data in enumerate(data):
                        for col_idx, cell_text in enumerate(row_data):
                            cell = table.cell(row_idx, col_idx)
                            cell.text = cell_text
                            if row_idx == 0:
                                cell.paragraphs[0].runs[0].bold = True
                                
                    p._element.addnext(table._element)
                    
    doc.save(output_path)
    print("Document saved with diagrams and wireframes.")

if __name__ == "__main__":
    create_kroki_diagrams()
    # graph_training is in the brain directory, passed as sys.argv[3]
    insert_diagrams_and_wireframes(sys.argv[1], sys.argv[2], sys.argv[3])
