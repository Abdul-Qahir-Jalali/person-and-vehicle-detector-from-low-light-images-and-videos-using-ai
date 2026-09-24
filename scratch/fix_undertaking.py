import sys
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def fix_undertaking(input_path, output_path):
    doc = Document(input_path)
    
    start_idx = -1
    
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip().lower()
        if txt == "undertaking":
            start_idx = i
            break
            
    if start_idx != -1:
        # Delete the badly formatted signature lines which should be the next few paragraphs
        # Specifically, look for "(student’s signature)"
        for j in range(start_idx, start_idx + 10):
            if j < len(doc.paragraphs):
                p = doc.paragraphs[j]
                if "(student’s signature)" in p.text.lower() or "abdul qahir" in p.text.lower() or "2021-umdb-001017" in p.text.lower():
                    # Delete this paragraph
                    p_element = p._element
                    parent = p_element.getparent()
                    if parent is not None:
                        parent.remove(p_element)
                        
        # Let's insert the table after the main text of the undertaking
        # The main text is probably at start_idx + 1 or start_idx + 2
        # We find the paragraph that contains "Where material has been used"
        target_p = None
        for j in range(start_idx, start_idx + 10):
            if j < len(doc.paragraphs):
                p = doc.paragraphs[j]
                if "where material has been used" in p.text.lower():
                    target_p = p
                    break
                    
        if target_p is not None:
            # We want to add a table AFTER this paragraph.
            table = doc.add_table(rows=1, cols=2)
            
            # Remove table borders (invisible table)
            # By default, a newly added table in python-docx with no style specified might have no borders,
            # but to be safe, we can just not add the tblBorders element.
            
            cell0 = table.cell(0, 0)
            p0 = cell0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.add_run("\n\n___________________________\n")
            p0.add_run("(Student’s Signature)\n")
            r0 = p0.add_run("Abdul Qahir Jalali\n")
            r0.bold = True
            p0.add_run("2021-UMDB-001017")
            
            cell1 = table.cell(0, 1)
            p1 = cell1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.add_run("\n\n___________________________\n")
            p1.add_run("(Student’s Signature)\n")
            r1 = p1.add_run("Raja Sharyar Muneer\n")
            r1.bold = True
            p1.add_run("2021-UMDB-001041")
            
            target_p._element.addnext(table._element)
            print("Undertaking table inserted.")
            
    doc.save(output_path)
    print("Undertaking Fixed and Saved")

if __name__ == "__main__":
    fix_undertaking(sys.argv[1], sys.argv[2])
