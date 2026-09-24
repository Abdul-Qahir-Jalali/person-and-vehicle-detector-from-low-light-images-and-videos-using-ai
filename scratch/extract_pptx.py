import sys
import os

try:
    from pptx import Presentation
except ImportError:
    print("Please run with: uv run --with python-pptx script.py")
    sys.exit(1)

def extract_pptx(path):
    prs = Presentation(path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(1)
    path = sys.argv[1]
    sys.stdout.reconfigure(encoding='utf-8')
    print(extract_pptx(path))
