import sys
try:
    from docx import Document
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def extract_figures(path):
    doc = Document(path)
    
    in_lof = False
    print("--- LIST OF FIGURES SECTION ---")
    for i, p in enumerate(doc.paragraphs[:250]):
        txt = p.text.strip().lower()
        if txt == "list of figures" and "heading" in p.style.name.lower():
            in_lof = True
            continue
            
        if in_lof:
            if p.style.name.lower() == "heading 1":
                in_lof = False
            else:
                print(f"[{i}] {p.style.name} | Text: {repr(p.text)}")
                
    print("\n--- ACTUAL FIGURES IN DOCUMENT ---")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith("Figure ") or txt.startswith("Table "):
            print(f"[{i}] {p.style.name}: {txt}")

if __name__ == "__main__":
    extract_figures(sys.argv[1])
