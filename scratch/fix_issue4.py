import sys
import re
try:
    from docx import Document
except ImportError:
    print("Please run with: uv run --with python-docx script.py")
    sys.exit(1)

def fix_issue4(input_path, output_path):
    doc = Document(input_path)
    
    inclusions_seen = False
    
    from docx.document import Document as _Document
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.table import _Cell, Table
    from docx.text.paragraph import Paragraph

    def iter_block_items(parent):
        if isinstance(parent, _Document):
            parent_elm = parent.element.body
        elif isinstance(parent, _Cell):
            parent_elm = parent._tc
        else:
            return

        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    paragraphs = []
    def extract_all_paragraphs(parent):
        for block in iter_block_items(parent):
            if isinstance(block, Paragraph):
                paragraphs.append(block)
            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        extract_all_paragraphs(cell)
                        
    extract_all_paragraphs(doc)
    
    last_text = ""
    for i, p in enumerate(paragraphs):
        txt = p.text.strip().lower()
        if not txt:
            continue
            
        # 1. Section 1.2.1 (Inclusions and Exclusions) duplication
        if txt.startswith("inclusions:"):
            if inclusions_seen and txt == last_text:
                if len(p.runs) > 0:
                    for run in p.runs:
                        run.text = ""
            else:
                inclusions_seen = True
                last_text = txt
        else:
            last_text = txt
            
        # Helper to safely insert bullets
        def insert_bullets(base_p, next_p, bullet_texts):
            for b_txt in bullet_texts:
                new_p = next_p.insert_paragraph_before("• " + b_txt)
                new_p.style = base_p.style
                
        # 2. Section 4.2.1 (OpenCV Integration)
        if "the process includes:" in txt:
            next_p = paragraphs[i+1] if i + 1 < len(paragraphs) else doc.add_paragraph()
            bullets = [
                "Dehazing: Applying Dark Channel Prior (DCP) algorithms to remove fog and improve visibility.",
                "Low-Light Enhancement: Utilizing Histogram Equalization to adjust image contrast and brighten dark areas.",
                "Noise Reduction: Applying Gaussian blur and bilateral filtering to smooth out rain streaks and artifacts."
            ]
            insert_bullets(p, next_p, bullets)
            
        # 3. Section 4.4.2 (Frontend Development)
        if "key features include:" in txt:
            next_p = paragraphs[i+1] if i + 1 < len(paragraphs) else doc.add_paragraph()
            bullets = [
                "A drag-and-drop file upload interface for images and videos (MP4, AVI, JPG, PNG).",
                "Real-time processing indicators to show enhancement progress.",
                "A dynamic result viewer that overlays bounding boxes and confidence scores directly onto the processed media."
            ]
            insert_bullets(p, next_p, bullets)

    doc.save(output_path)
    print("Issue 4 Fixed and Saved")

if __name__ == "__main__":
    fix_issue4(sys.argv[1], sys.argv[2])
